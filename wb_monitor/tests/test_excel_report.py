"""Тесты отчёта: структура строго повторяет sample_report.xlsx."""

from __future__ import annotations

from datetime import date, datetime

import openpyxl
from sqlalchemy import select

from wb_monitor.models import Client, PriceSnapshot, Tier, TrackedSKU
from wb_monitor.reports.excel_report import (
    ReportData,
    ReportRow,
    build_report_file,
    collect_report_data,
    format_period,
)
from wb_monitor.scheduler.jobs import weekly_report_job
from wb_monitor.tests.conftest import FakeBot

PERIOD = (date(2026, 7, 20), date(2026, 7, 26))


def _sample_data() -> ReportData:
    return ReportData(
        client_name="ИП Тестов",
        period_start=PERIOD[0],
        period_end=PERIOD[1],
        rows=[
            ReportRow("Носки свои", "18234561", True, "Носки", 590, 590, 42),
            ReportRow("Носки конкурента", "22190345", False, "Носки", 450, 369, 120),
            ReportRow(
                "Термобельё конкурента", "21094821", False, "Термо", 1450, 1450, 0,
                stock_zeroed_at=date(2026, 7, 24),
            ),
            ReportRow("Шапка конкурента", "20456781", False, "Аксессуары", 480, 510, 60),
            ReportRow("Шарф конкурента", "21567890", False, "Аксессуары", 610, 610, 5),
        ],
        alerts_sent=3,
    )


def test_report_structure_matches_sample(tmp_path):
    path = build_report_file(_sample_data(), tmp_path / "r.xlsx")
    wb = openpyxl.load_workbook(path)

    assert wb.sheetnames == ["Сводка", "Цены и остатки"]

    ws = wb["Сводка"]
    assert ws["B2"].value == "Еженедельный отчёт: мониторинг конкурентов на Wildberries"
    assert ws["B3"].value == "Клиент: ИП Тестов"
    assert ws["B4"].value == "Период: 20-26 июля 2026"
    assert (ws["B6"].value, ws["D6"].value) == ("Товаров под наблюдением", 5)
    assert (ws["B7"].value, ws["D7"].value) == ("из них ваши", 1)
    assert (ws["B8"].value, ws["D8"].value) == ("из них конкуренты", 4)
    assert (ws["B9"].value, ws["D9"].value) == ("Алертов отправлено за неделю", 3)
    assert (ws["B10"].value, ws["D10"].value) == ("Обнулений остатка у конкурентов", 1)
    assert ws["B12"].value == "Главное за неделю"
    assert "снизил цену на 18%" in ws["B13"].value
    assert "обнулил остаток 24 июля" in str(ws["B13"].value) + str(ws["B14"].value)

    ws = wb["Цены и остатки"]
    assert [c.value for c in ws[2]] == [
        "Товар", "Артикул (nmId)", "Тип", "Категория",
        "Цена 20.07", "Цена 26.07", "Изменение", "Остаток сейчас", "Статус",
    ]
    assert ws.freeze_panes == "A3"
    # формула изменения — формула Excel, не число
    assert ws["G3"].value == "=IFERROR((F3-E3)/E3,0)"
    assert ws["G7"].value == "=IFERROR((F7-E7)/E7,0)"
    assert ws["G3"].number_format == r"\+0.0%;\-0.0%;0.0%"
    assert ws["E3"].number_format == '#,##0" руб."'
    # шапка: белый жирный на синей заливке
    assert ws["A2"].font.b and ws["A2"].fill.fgColor.rgb == "002F5597"


def test_report_statuses(tmp_path):
    path = build_report_file(_sample_data(), tmp_path / "r.xlsx")
    ws = openpyxl.load_workbook(path)["Цены и остатки"]
    statuses = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=9).value
                for r in range(3, 8)}
    assert statuses["Носки свои"] == "OK"
    assert statuses["Носки конкурента"] == "Цена упала"        # -18%
    assert statuses["Термобельё конкурента"] == "Обнулился остаток"
    assert statuses["Шапка конкурента"] == "Цена выросла"      # +6.25%
    assert statuses["Шарф конкурента"] == "Мало остатка"       # 5 шт.


def test_format_period_cross_month():
    assert format_period(date(2026, 7, 28), date(2026, 8, 3)) == "28 июля - 3 августа 2026"


def test_quiet_week_highlight(tmp_path):
    data = ReportData("К", *PERIOD, rows=[ReportRow("Т", "1", True, "", 100, 100, 9)])
    path = build_report_file(data, tmp_path / "r.xlsx")
    ws = openpyxl.load_workbook(path)["Сводка"]
    assert "спокойная" in ws["B13"].value


async def test_weekly_report_job_end_to_end(session_factory, tmp_path):
    """Ручной сценарий: клиент + снепшоты за неделю -> файл + отправка ботом."""
    async with session_factory() as session:
        client = Client(name="ИП Е2Е", telegram_chat_id=333, tier=Tier.BASIC.value, price_rub=990)
        session.add(client)
        await session.flush()
        sku = TrackedSKU(
            client_id=client.id, wb_article_id="777", display_name="Товар", category="Носки"
        )
        session.add(sku)
        await session.flush()
        session.add_all([
            PriceSnapshot(sku_id=sku.id, checked_at=datetime(2026, 7, 20, 8), price=500, stock_qty=10),
            PriceSnapshot(sku_id=sku.id, checked_at=datetime(2026, 7, 26, 8), price=430, stock_qty=0),
        ])
        await session.commit()
        client_id = client.id

    bot = FakeBot()
    paths = await weekly_report_job(
        session_factory, bot, str(tmp_path), period_end=date(2026, 7, 26)
    )
    assert len(paths) == 1
    assert len(bot.messages) == 1  # документ ушёл в чат клиента

    ws = openpyxl.load_workbook(paths[0])["Цены и остатки"]
    assert ws["A3"].value == "Товар"
    assert (ws["E3"].value, ws["F3"].value, ws["H3"].value) == (500, 430, 0)
    assert ws["I3"].value == "Обнулился остаток"
