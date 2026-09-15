"""Генерация еженедельного Excel-отчёта. Структура = sample_report.xlsx:

лист «Сводка»          — шапка, счётчики, «Главное за неделю»
лист «Цены и остатки»  — таблица по SKU, заголовки в строке 2, freeze A3,
                         колонка G «Изменение» — формула Excel
                         =IFERROR((Fn-En)/En,0), формат +0.0%/-0.0%
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from wb_monitor.models import Alert, Client, PriceSnapshot, TrackedSKU
from wb_monitor.scheduler import thresholds

# --- статусы в колонке I (те же формулировки, что в образце) ---
STATUS_OK = "OK"
STATUS_PRICE_DOWN = "Цена упала"
STATUS_PRICE_UP = "Цена выросла"
STATUS_STOCK_ZERO = "Обнулился остаток"
STATUS_LOW_STOCK = "Мало остатка"

_MONTHS_GEN = [
    "января", "февраля", "марта", "апреля", "мая", "июня",
    "июля", "августа", "сентября", "октября", "ноября", "декабря",
]

_HEADER_FILL = PatternFill("solid", fgColor="2F5597")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_THIN = Side(style="thin")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_PRICE_FMT = '#,##0" руб."'
_PCT_FMT = r"\+0.0%;\-0.0%;0.0%"
_STATUS_FONTS = {
    STATUS_PRICE_DOWN: Font(bold=True, color="C00000"),
    STATUS_STOCK_ZERO: Font(bold=True, color="C00000"),
    STATUS_PRICE_UP: Font(color="2F5597"),
}


@dataclass
class ReportRow:
    name: str
    nm_id: str
    is_own: bool
    category: str
    price_start: float | None   # эффективная цена на начало периода
    price_end: float | None     # ... на конец периода
    stock_now: int | None
    stock_zeroed_at: date | None = None  # дата обнуления в течение недели, если было


@dataclass
class ReportData:
    client_name: str
    period_start: date
    period_end: date
    rows: list[ReportRow] = field(default_factory=list)
    alerts_sent: int = 0

    @property
    def own_count(self) -> int:
        return sum(r.is_own for r in self.rows)

    @property
    def competitor_zeroed_count(self) -> int:
        return sum(1 for r in self.rows if not r.is_own and r.stock_zeroed_at is not None)


def format_period(start: date, end: date) -> str:
    """«20-26 июля 2026» или «28 июля - 3 августа 2026»."""
    if start.month == end.month:
        return f"{start.day}-{end.day} {_MONTHS_GEN[end.month - 1]} {end.year}"
    return (
        f"{start.day} {_MONTHS_GEN[start.month - 1]} - "
        f"{end.day} {_MONTHS_GEN[end.month - 1]} {end.year}"
    )


# ------------------------------------------------------------ сбор данных


async def collect_report_data(
    session: AsyncSession, client: Client, period_start: date, period_end: date
) -> ReportData:
    """Снепшоты за период по всем активным SKU клиента -> данные отчёта.

    Цена «на начало» — первый снепшот в окне (если его нет — последний до
    окна, чтобы неделя без изменений не выглядела пустой), «на конец» —
    последний в окне.
    """
    window_start = datetime.combine(period_start, time.min)
    window_end = datetime.combine(period_end, time.max)

    data = ReportData(
        client_name=client.name, period_start=period_start, period_end=period_end
    )

    skus = (
        await session.execute(
            select(TrackedSKU)
            .where(TrackedSKU.client_id == client.id, TrackedSKU.is_active)
            .order_by(TrackedSKU.category, TrackedSKU.id)
        )
    ).scalars().all()

    for sku in skus:
        in_window = (
            await session.execute(
                select(PriceSnapshot)
                .where(
                    PriceSnapshot.sku_id == sku.id,
                    PriceSnapshot.checked_at >= window_start,
                    PriceSnapshot.checked_at <= window_end,
                )
                .order_by(PriceSnapshot.checked_at, PriceSnapshot.id)
            )
        ).scalars().all()

        first = in_window[0] if in_window else None
        last = in_window[-1] if in_window else None
        if first is None:
            first = (
                await session.execute(
                    select(PriceSnapshot)
                    .where(
                        PriceSnapshot.sku_id == sku.id,
                        PriceSnapshot.checked_at < window_start,
                    )
                    .order_by(PriceSnapshot.checked_at.desc(), PriceSnapshot.id.desc())
                    .limit(1)
                )
            ).scalar_one_or_none()

        zeroed_at: date | None = None
        prev_stock: int | None = None
        for snap in in_window:
            if prev_stock is not None and prev_stock > 0 and snap.stock_qty == 0:
                zeroed_at = snap.checked_at.date()
            prev_stock = snap.stock_qty

        data.rows.append(ReportRow(
            name=sku.display_name,
            nm_id=sku.wb_article_id,
            is_own=sku.is_own_product,
            category=sku.category or "",
            price_start=first.effective_price if first else None,
            price_end=last.effective_price if last else None,
            stock_now=last.stock_qty if last else None,
            stock_zeroed_at=zeroed_at,
        ))

    data.alerts_sent = (
        await session.execute(
            select(func.count())
            .select_from(Alert)
            .where(
                Alert.client_id == client.id,
                Alert.sent_at >= window_start,
                Alert.sent_at <= window_end,
            )
        )
    ).scalar_one()
    return data


# ------------------------------------------------------------- генерация


def _row_status(row: ReportRow) -> str:
    if row.stock_now == 0:
        return STATUS_STOCK_ZERO
    if row.price_start and row.price_end:
        change = (row.price_end - row.price_start) / row.price_start
        if change <= -thresholds.REPORT_PRICE_CHANGE_THRESHOLD:
            return STATUS_PRICE_DOWN
        if change >= thresholds.REPORT_PRICE_CHANGE_THRESHOLD:
            return STATUS_PRICE_UP
    if row.stock_now is not None and 0 < row.stock_now <= thresholds.LOW_STOCK_QTY:
        return STATUS_LOW_STOCK
    return STATUS_OK


def _highlights(data: ReportData) -> list[str]:
    """Строки «Главного за неделю» — по тем же событиям, что в образце."""
    bullets: list[str] = []

    drops: list[tuple[float, ReportRow]] = []
    for row in data.rows:
        if row.is_own or not row.price_start or not row.price_end:
            continue
        change = (row.price_end - row.price_start) / row.price_start
        if change <= -thresholds.REPORT_PRICE_CHANGE_THRESHOLD:
            drops.append((change, row))
    for change, row in sorted(drops)[:2]:
        bullets.append(
            f"- Конкурент «{row.name}» снизил цену на {-change:.0%} "
            f"(с {row.price_start:.0f} до {row.price_end:.0f} руб.) - "
            f"стоит проверить свою цену на этот SKU."
        )

    for row in data.rows:
        if not row.is_own and row.stock_zeroed_at is not None:
            day = row.stock_zeroed_at
            bullets.append(
                f"- Конкурент «{row.name}» обнулил остаток "
                f"{day.day} {_MONTHS_GEN[day.month - 1]} - "
                f"окно возможностей для рекламы вашей позиции."
            )

    if not bullets:
        bullets.append(
            "- Неделя спокойная: без резких изменений цен и остатков у конкурентов."
        )
    return bullets[:4]


def build_report_file(data: ReportData, out_path: str | Path) -> Path:
    """Собрать .xlsx по структуре sample_report.xlsx и вернуть путь."""
    wb = Workbook()
    _build_summary_sheet(wb.active, data)
    _build_table_sheet(wb.create_sheet(), data)
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


def _build_summary_sheet(ws: Worksheet, data: ReportData) -> None:
    ws.title = "Сводка"
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["H"].width = 3

    ws["B2"] = "Еженедельный отчёт: мониторинг конкурентов на Wildberries"
    ws["B2"].font = Font(bold=True, size=14)
    ws["B3"] = f"Клиент: {data.client_name}"
    ws["B4"] = f"Период: {format_period(data.period_start, data.period_end)}"

    counters = [
        ("Товаров под наблюдением", len(data.rows)),
        ("из них ваши", data.own_count),
        ("из них конкуренты", len(data.rows) - data.own_count),
        ("Алертов отправлено за неделю", data.alerts_sent),
        ("Обнулений остатка у конкурентов", data.competitor_zeroed_count),
    ]
    for i, (label, value) in enumerate(counters):
        row = 6 + i
        ws.cell(row=row, column=2, value=label).font = Font(bold=True)
        ws.cell(row=row, column=4, value=value)

    ws["B12"] = "Главное за неделю"
    ws["B12"].font = Font(bold=True)
    for i, text in enumerate(_highlights(data)):
        ws.cell(row=13 + i, column=2, value=text)
        ws.row_dimensions[13 + i].height = 30


def _build_table_sheet(ws: Worksheet, data: ReportData) -> None:
    ws.title = "Цены и остатки"
    for col, width in {"A": 30, "B": 15, "C": 12, "D": 14, "E": 12, "H": 14, "I": 18}.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A3"

    headers = [
        "Товар", "Артикул (nmId)", "Тип", "Категория",
        f"Цена {data.period_start:%d.%m}", f"Цена {data.period_end:%d.%m}",
        "Изменение", "Остаток сейчас", "Статус",
    ]
    ws.row_dimensions[2].height = 26.85
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=2, column=col, value=header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _BORDER

    for i, row in enumerate(data.rows):
        r = 3 + i
        status = _row_status(row)
        values = [
            row.name,
            row.nm_id,
            "Свой" if row.is_own else "Конкурент",
            row.category,
            row.price_start,
            row.price_end,
            f"=IFERROR((F{r}-E{r})/E{r},0)",  # формула Excel, не питоновское число
            row.stock_now,
            status,
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            cell.border = _BORDER
            if col in (5, 6):
                cell.number_format = _PRICE_FMT
            elif col == 7:
                cell.number_format = _PCT_FMT
        status_font = _STATUS_FONTS.get(status)
        if status_font is not None:
            ws.cell(row=r, column=9).font = status_font
